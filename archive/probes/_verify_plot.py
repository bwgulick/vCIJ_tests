import openpyxl
from openpyxl.utils import get_column_letter
wf=openpyxl.load_workbook("Cij analysis/Vanadium_Cij_Brian_Claude.xlsx", data_only=False)
wv=openpyxl.load_workbook("Cij analysis/Vanadium_Cij_Brian_Claude.xlsx", data_only=True)
assert "Plot Data" in wf.sheetnames, "sheet missing"
sf,sv=wf["Plot Data"],wv["Plot Data"]
print("sheets:", wf.sheetnames)
# headers row2
hdr=[sf.cell(2,c).value for c in range(1,18)]
print("HDRS:", " | ".join(str(h) for h in hdr))
# formula check row3 (first data) and value check rows 3,4,24
print("\nrow3 formulas:")
for c in range(1,18):
    print(f"  {get_column_letter(c)}3 = {sf.cell(3,c).value}")
print("\nvalues (out row: P, Vp100 sig, C11 sig, C12 sig):")
for r in [3,4,24]:
    vals=[sv.cell(r,c).value for c in range(1,18)]
    def f(v): return f"{v:.4f}" if isinstance(v,(int,float)) else str(v)
    print(f"  r{r}: P={f(vals[0])} L={f(vals[1])}±{f(vals[2])} Vp100={f(vals[3])}±{f(vals[4])} Vs100={f(vals[5])}±{f(vals[6])} Vp111={f(vals[7])}±{f(vals[8])} Vs111={f(vals[9])}±{f(vals[10])} C11={f(vals[11])}±{f(vals[12])} C44={f(vals[13])}±{f(vals[14])} C12={f(vals[15])}±{f(vals[16])}")
# check for any #REF! or None in data block rows 3..24
bad=[]
for r in range(3,25):
    for c in range(1,18):
        v=sv.cell(r,c).value
        if v is None or (isinstance(v,str) and "REF" in v): bad.append(f"{get_column_letter(c)}{r}={v}")
print("\nBAD cells in data block:", bad if bad else "NONE")
