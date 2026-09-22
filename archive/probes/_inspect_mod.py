import openpyxl
wb = openpyxl.load_workbook("Cij analysis/Vanadium_Cij_Brian_Claude.xlsx", data_only=False)
ws = wb["No Vp 100"]

def colletters(idx):
    from openpyxl.utils import get_column_letter
    return get_column_letter(idx)

# Header row 1 for the modulus columns and covariance block
print("=== Headers row 1 for CB..CI ===")
for col in ["CB","CC","CD","CE","CF","CG","CH","CI"]:
    print(col, "hdr:", repr(ws[col+"1"].value), "| r2 formula:", repr(ws[col+"2"].value))

print()
print("=== Cij value cols AJ/AK/AL row1 header + r2 formula ===")
for col in ["AJ","AK","AL"]:
    print(col, "hdr:", repr(ws[col+"1"].value), "| r2:", repr(ws[col+"2"].value))

print()
print("=== Covariance / sigma block ET..FO row1 headers ===")
from openpyxl.utils import column_index_from_string, get_column_letter
for c in range(column_index_from_string("ET"), column_index_from_string("FO")+1):
    L = get_column_letter(c)
    print(L, "hdr:", repr(ws[L+"1"].value), "| r2:", repr(ws[L+"2"].value))
