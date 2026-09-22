import openpyxl
from openpyxl.utils import get_column_letter, column_index_from_string
fp = "Cij analysis/Vanadium_Cij_Brian_Claude.xlsx"
wf = openpyxl.load_workbook(fp, data_only=False)
wv = openpyxl.load_workbook(fp, data_only=True)
sn="No Vs 111"; sf,sv=wf[sn],wv[sn]
# pressure column A and length rows 2..25
print("row | A(P) | M(len) | Q(sigL) | AG(C11) | AH(C44) | AI(C12)")
for r in range(2,26):
    def V(col): return sv[f"{col}{r}"].value
    print(f"{r:>3} | {str(V('A'))[:8]:<8} | {str(V('M'))[:7]:<7} | {str(V('Q'))[:7]:<7} | {str(V('AG'))[:7]:<7} | {str(V('AH'))[:7]:<7} | {str(V('AI'))[:7]}")
# ET..FO headers (cols 150..171) row1 and row2
print("\nET..FO block (row1 header / row2 formula / row2 value):")
for c in range(column_index_from_string("ET"), column_index_from_string("FO")+1):
    L=get_column_letter(c)
    h1=sf.cell(1,c).value; f2=sf.cell(2,c).value; v2=sv.cell(2,c).value
    print(f"{L:>3} h1={str(h1)[:26]:<26} f2={str(f2)[:30]:<30} v2={str(v2)[:14]}")
