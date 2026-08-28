import pandas as pd, numpy as np
from openpyxl.utils import get_column_letter
path = r"C:\Users\bgulick\Desktop\V_single_figures\VCIJplotdata.xlsx"

xl = pd.ExcelFile(path)
print("sheets:", xl.sheet_names)

# read with no header so we see the raw grid + can locate label rows/cols
raw = pd.read_excel(path, sheet_name="uncertainty calcs", header=None)
print("shape:", raw.shape)

# Dump columns ~ET-EU region (index 145..170) : show first ~6 rows as labels
lo, hi = 140, 172
print(f"\n=== columns {get_column_letter(lo+1)}({lo})..{get_column_letter(hi+1)}({hi}), first 6 rows ===")
for c in range(lo, min(hi, raw.shape[1])):
    col = raw.iloc[0:6, c].tolist()
    # only print columns that have any non-nan
    if any(pd.notna(v) for v in col):
        print(f"col {c} ({get_column_letter(c+1)}): {col}")
