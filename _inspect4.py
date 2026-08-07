import openpyxl
from openpyxl.utils import get_column_letter, column_index_from_string
fp = "Cij analysis/Vanadium_Cij_Brian_Claude.xlsx"
wf = openpyxl.load_workbook(fp, data_only=False)
wv = openpyxl.load_workbook(fp, data_only=True)
sn="No Vp 100"; sf,sv=wf[sn],wv[sn]
# ET block full formulas
print("-- ET..FO formulas (No Vp 100) --")
for c in range(column_index_from_string("ET"), column_index_from_string("FO")+1):
    L=get_column_letter(c)
    print(f"{L:>3} {sf.cell(1,c).value} :: {sf.cell(2,c).value}")
# find experimental Cij cells: scan AJ..DZ for headers containing C11/C44/C12 exp/calc
print("\n-- headers AJ..DZ mentioning C11/C44/C12/exp/calc --")
for c in range(column_index_from_string("AJ"), column_index_from_string("DZ")+1):
    h=sf.cell(1,c).value
    if h and any(k in str(h) for k in ["C11","C44","C12","exp","calc","C 11","C 44","C 12"]):
        L=get_column_letter(c)
        print(f"{L:>3} h={str(h)[:26]:<26} f2={str(sf.cell(2,c).value)[:34]} v2={str(sv.cell(2,c).value)[:12]}")
